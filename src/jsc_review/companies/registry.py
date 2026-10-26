"""Load the companies-to-audit directory from the local Excel file.

The source Excel groups rows by sector with a banner row like ``البنوك``
followed by the header row ``اسم الشركة, الاسم المختصر, ...`` and then the
real company rows. The parser:

* ignores banner rows (any row with only a single non-empty cell)
* ignores header rows (they match the known header tuple)
* treats a row as a company iff it has a non-empty name AND a non-empty
  numeric code or ticker
* collapses duplicates across sheets (the Excel ships both an abridged
  worksheet and the full listing)

The parsed registry is written as JSON at
``Settings.companies_registry_path`` so the API does not have to open the
Excel on every request. ``load_registry`` reads that JSON. ``build_registry``
takes an Excel path and writes the JSON cache.
"""

from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from pydantic import BaseModel, ConfigDict, Field

# The canonical column names as they appear in the Excel header row.
_HEADER_CELLS: tuple[str, ...] = (
    "اسم الشركة",
    "الاسم المختصر للشركة",
    "الرمز الحرفي",
    "الرمز الرقمي",
    "السوق",
    "الأسهم المدرجة",
    "الحالة",
    "ملاحظات",
)


class Company(BaseModel):
    """One listed company we may audit."""

    model_config = ConfigDict(extra="forbid")

    name_ar: str
    short_name_ar: str = ""
    ticker: str = ""
    numeric_code: str = ""
    market: str = ""
    listed_shares: int | None = None
    sector_ar: str = ""
    status: str = ""
    notes: str = ""

    @property
    def key(self) -> str:
        """Return the stable identifier the mapping store uses.

        The ticker is preferred because it is short and unique. The numeric
        code is a safe fallback for the handful of rows without a ticker in
        the source Excel. When both are empty, the Arabic name is used as a
        last-resort key.
        """

        if self.ticker:
            return self.ticker
        if self.numeric_code:
            return self.numeric_code
        return self.name_ar


class CompanyRegistryPayload(BaseModel):
    """Serializable envelope written to the on-disk cache."""

    model_config = ConfigDict(extra="forbid")

    source_excel: str = ""
    companies: list[Company] = Field(default_factory=list)


@dataclass(frozen=True)
class CompanyRegistry:
    """In-memory lookup over companies by ticker / numeric_code / name."""

    companies: tuple[Company, ...]

    def by_key(self, key: str) -> Company | None:
        for company in self.companies:
            if company.key == key:
                return company
        return None

    def by_ticker(self, ticker: str) -> Company | None:
        if not ticker:
            return None
        folded = ticker.strip().upper()
        for company in self.companies:
            if company.ticker.strip().upper() == folded:
                return company
        return None

    def by_numeric_code(self, code: str) -> Company | None:
        if not code:
            return None
        folded = code.strip()
        for company in self.companies:
            if company.numeric_code.strip() == folded:
                return company
        return None


def _is_header_row(row: tuple[Any, ...]) -> bool:
    """Return True when the row is the Arabic column-header row."""

    cells = tuple((str(cell).strip() if cell is not None else "") for cell in row)
    non_empty = tuple(cell for cell in cells if cell)
    if not non_empty:
        return True  # treat empty rows as headers so parse_rows skips them
    return non_empty[: len(_HEADER_CELLS)] == _HEADER_CELLS[: len(non_empty)]


def _is_sector_banner(row: tuple[Any, ...]) -> tuple[bool, str]:
    """Detect sector banner rows and return (is_banner, sector_name)."""

    non_empty = [
        (position, str(cell).strip())
        for position, cell in enumerate(row)
        if cell is not None and str(cell).strip()
    ]
    if len(non_empty) == 1 and non_empty[0][0] == 0:
        return True, non_empty[0][1]
    return False, ""


def _as_int(value: Any) -> int | None:
    """Coerce a cell to int when possible, otherwise ``None``."""

    if value is None:
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        text = str(value).strip().replace(",", "")
        if not text:
            return None
        try:
            return int(float(text))
        except (TypeError, ValueError):
            return None


def _as_str(value: Any) -> str:
    """Return the cell as a stripped string. Numeric cells are coerced sanely."""

    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_company_rows(rows: Iterable[tuple[Any, ...]]) -> list[Company]:
    """Return every company found in the raw rows of one sheet.

    Sector banners update the current sector label applied to the next rows.
    Header rows are skipped. Rows with no name are skipped. The order of
    appearance is preserved.
    """

    companies: list[Company] = []
    sector = ""
    for row in rows:
        banner, name = _is_sector_banner(row)
        if banner:
            sector = name
            continue
        if _is_header_row(row):
            continue
        name_ar = _as_str(row[0] if len(row) > 0 else "")
        if not name_ar:
            continue
        short_name = _as_str(row[1] if len(row) > 1 else "")
        ticker = _as_str(row[2] if len(row) > 2 else "")
        numeric_code = _as_str(row[3] if len(row) > 3 else "")
        if not ticker and not numeric_code:
            # Looks like another banner that happens to span one column but
            # not at position 0, or a stray title row. Skip it.
            continue
        companies.append(
            Company(
                name_ar=name_ar,
                short_name_ar=short_name,
                ticker=ticker,
                numeric_code=numeric_code,
                market=_as_str(row[4] if len(row) > 4 else ""),
                listed_shares=_as_int(row[5] if len(row) > 5 else None),
                sector_ar=sector,
                status=_as_str(row[6] if len(row) > 6 else ""),
                notes=_as_str(row[7] if len(row) > 7 else ""),
            )
        )
    return companies


def _deduplicate_companies(companies: Iterable[Company]) -> list[Company]:
    """Collapse duplicate companies by stable key, keeping the fullest record.

    When the same key appears twice (e.g. once in the abridged worksheet and
    once in the full one), the row with more populated fields wins. Ties are
    broken by first appearance so output order is deterministic.
    """

    seen: dict[str, Company] = {}
    order: list[str] = []
    for company in companies:
        key = company.key
        if key not in seen:
            seen[key] = company
            order.append(key)
            continue
        if _completeness_score(company) > _completeness_score(seen[key]):
            seen[key] = company
    return [seen[key] for key in order]


def _completeness_score(company: Company) -> int:
    score = 0
    for field in (
        company.name_ar,
        company.short_name_ar,
        company.ticker,
        company.numeric_code,
        company.market,
        company.sector_ar,
        company.status,
        company.notes,
    ):
        if field:
            score += 1
    if company.listed_shares is not None:
        score += 1
    return score


def companies_from_excel(path: Path) -> list[Company]:
    """Return every company across every sheet of the Excel file."""

    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)
    collected: list[Company] = []
    try:
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            collected.extend(parse_company_rows(sheet.iter_rows(values_only=True)))
    finally:
        workbook.close()
    return _deduplicate_companies(collected)


def build_registry(
    source_excel: Path, cache_path: Path
) -> CompanyRegistryPayload:
    """Parse the Excel and write the JSON cache. Returns the written payload."""

    companies = companies_from_excel(source_excel)
    payload = CompanyRegistryPayload(
        source_excel=str(source_excel),
        companies=companies,
    )
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = cache_path.with_suffix(cache_path.suffix + ".tmp")
    tmp.write_text(payload.model_dump_json(indent=2), encoding="utf-8")
    tmp.replace(cache_path)
    return payload


def load_registry(cache_path: Path) -> CompanyRegistry:
    """Read the JSON cache. Returns an empty registry when the file is missing."""

    if not cache_path.exists():
        return CompanyRegistry(companies=())
    payload = CompanyRegistryPayload.model_validate_json(
        cache_path.read_text(encoding="utf-8")
    )
    return CompanyRegistry(companies=tuple(payload.companies))


def load_registry_payload(cache_path: Path) -> CompanyRegistryPayload:
    """Read the raw cache payload (companies + source metadata)."""

    if not cache_path.exists():
        return CompanyRegistryPayload()
    return CompanyRegistryPayload.model_validate_json(
        cache_path.read_text(encoding="utf-8")
    )


def to_json_payload(registry: CompanyRegistry, source_excel: str = "") -> str:
    """Serialize an in-memory registry to the on-disk JSON shape."""

    return CompanyRegistryPayload(
        source_excel=source_excel, companies=list(registry.companies)
    ).model_dump_json(indent=2)
